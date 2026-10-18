#!/usr/bin/env python3
"""Xuất file BA (.xlsx tải từ Google Sheets: Tệp ▸ Tải xuống ▸ Microsoft Excel) thành fixture JSON cho
worker/test/realdata.mjs và tools/sim_week.mjs.

  python3 tools/ba_fixture.py <BA.xlsx> [.local/fixture_ba.json]

Fixture chứa DỮ LIỆU THẬT (tên, SĐT) → chỉ để trong .local/ (đã .gitignore), không in, không đẩy lên repo.
Mã hoá ô: ngày → {"$d": "yyyy-mm-dd"} · giờ → {"$t": "HH:MM"} · số / chữ giữ nguyên. Cần: pip install openpyxl
"""
import datetime, json, os, sys
import openpyxl

TABS = ['MEMBERS', 'SESSION LOG', 'TRIAL', 'B0B']

def cell(v):
    if isinstance(v, datetime.datetime):
        if v.hour or v.minute or v.second:
            if v.year < 1901: return {'$t': v.strftime('%H:%M')}
        return {'$d': v.strftime('%Y-%m-%d')}
    if isinstance(v, datetime.date): return {'$d': v.isoformat()}
    if isinstance(v, datetime.time): return {'$t': v.strftime('%H:%M')}
    if v is None: return ''
    return v

def main():
    if len(sys.argv) < 2: sys.exit(__doc__)
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(__file__), '..', '.local', 'fixture_ba.json')
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    names = list(wb.sheetnames)
    pick = [t for t in TABS if t in names] + [n for n in names if '⚙' in n and n.replace(' ', '').endswith('2')]
    fx = {}
    for n in pick:
        rows = []
        for r in wb[n].iter_rows(values_only=True):
            r = [cell(v) for v in r]
            while r and r[-1] == '': r.pop()
            rows.append(r)
        while rows and not rows[-1]: rows.pop()
        fx[n] = rows
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    json.dump(fx, open(out, 'w', encoding='utf-8'), ensure_ascii=False)
    print('ghi', out, '·', ', '.join(f'{k}: {len(v)} hàng' for k, v in fx.items()))

if __name__ == '__main__':
    main()
